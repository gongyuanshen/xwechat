from __future__ import annotations
from .diagnostics import observed, event as diagnostic_event, executor_call
import logging

import asyncio
import base64
import hashlib
import io
import json
import time
import zipfile
import re
import math
import threading
from xml.etree import ElementTree
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor

from PIL import Image
from .agent_budget import pieces, size


# PDFium forbids simultaneous calls from different threads, even for separate
# documents. Keep the native lock out of generator yields and model inference.
_PDFIUM_LOCK = threading.Lock()


def image_url(data):
    with Image.open(io.BytesIO(data)) as img:
        img.seek(0)
        img = img.convert("RGB")
        img.thumbnail((1600, 1600))
        buffer = io.BytesIO()
        img.save(buffer, format="JPEG", quality=85)
    return "data:image/jpeg;base64," + base64.b64encode(buffer.getvalue()).decode()


def iter_document(data: bytes, suffix: str, *, include_images=True, scan_only=False, raw_images=False):
    """返回带位置的片段；不运行宏，不提取压缩包中的任意文件到磁盘。"""
    stream = io.BytesIO(data)
    if suffix in {".txt", ".md", ".csv"}:
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("gb18030")
        yield {"label": "正文", "text": text}
        return
    if suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ValueError("PDF 已加密")
        import pypdfium2 as pdfium
        document = None
        try:
            for i, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                yield {"label": f"第 {i + 1} 页", "text": text}
                # 无文字页面交给视觉模型；有文字页面的图片也参与理解。
                if not text.strip():
                    if not include_images:
                        yield {"label": f"第 {i + 1} 页扫描图", "skipped_image": True}
                        continue
                    with _PDFIUM_LOCK:
                        document = document or pdfium.PdfDocument(data)
                        pdf_page = document[i]
                        try:
                            scale = 2 if raw_images else 1.5
                            width, height = pdf_page.get_size()
                            if math.ceil(width * scale) * math.ceil(height * scale) > Image.MAX_IMAGE_PIXELS:
                                raise ValueError(f'第 {i + 1} 页渲染后超过本地图像像素上限，请先缩小或分割该页面。')
                            bitmap = pdf_page.render(scale=scale)
                            try:
                                buf = io.BytesIO()
                                bitmap.to_pil().save(buf, "PNG")
                                rendered = buf.getvalue()
                            finally:
                                bitmap.close()
                        finally:
                            pdf_page.close()
                    yield {"label": f"第 {i + 1} 页扫描图", "image": rendered if raw_images else image_url(rendered)}
                else:
                    if not include_images or scan_only:
                        for name in page.images.keys():
                            yield {"label": f"第 {i + 1} 页图片 {name}", "skipped_image": True}
                        continue
                    for n, img in enumerate(page.images):
                        yield {"label": f"第 {i + 1} 页图片 {n + 1}", "image": img.data if raw_images else image_url(img.data)}
        finally:
            if document is not None:
                with _PDFIUM_LOCK:
                    document.close()
        return
    if suffix not in {".docx", ".xlsx", ".pptx"}:
        raise ValueError("暂不支持此附件格式")
    with zipfile.ZipFile(stream) as archive:
        if sum(x.file_size for x in archive.infolist()) > 512 * 1024 * 1024:
            raise ValueError("附件解压后超过 512 MB，请拆分文件")
        # Office 对象库会预加载图片部件；纯文字模式直接读取 XML，避免解压图片数据。
        if not include_images and suffix in {'.docx', '.pptx'}:
            names = archive.namelist()
            if suffix == '.docx':
                ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
                root = ElementTree.fromstring(archive.read('word/document.xml'))
                body = root.find('w:body', ns)
                for i, block in enumerate(body if body is not None else []):
                    if block.tag.endswith('}tbl'):
                        rows = [' | '.join(''.join(cell.itertext()) for cell in row.findall('w:tc', ns)) for row in block.findall('w:tr', ns)]
                        yield {'label': f'表格 {i + 1}', 'text': '\n'.join(rows)}
                    else:
                        yield {'label': f'段落 {i + 1}', 'text': ''.join(node.text or '' for node in block.findall('.//w:t', ns))}
            else:
                ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
                slides = sorted((name for name in names if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)), key=lambda name: int(re.search(r'slide(\d+)\.xml', name)[1]))
                for i, name in enumerate(slides):
                    root = ElementTree.fromstring(archive.read(name))
                    yield {'label': f'幻灯片 {i + 1}', 'text': '\n'.join(''.join(t.text or '' for t in p.findall('.//a:t', ns)) for p in root.findall('.//a:p', ns))}
            for name in names:
                if '/media/' in name:
                    yield {'label': f'嵌入媒体 {Path(name).name}', 'skipped_image': True}
            return
    stream.seek(0)
    if suffix == ".docx":
        from docx import Document
        document = Document(stream)
        for i, para in enumerate(document.paragraphs):
            yield {"label": f"段落 {i + 1}", "text": para.text}
        for i, table in enumerate(document.tables):
            yield {"label": f"表格 {i + 1}", "text": "\n".join(" | ".join(c.text for c in row.cells) for row in table.rows)}
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        book = load_workbook(stream, read_only=True, data_only=False)
        try:
            for sheet in book:
                for i, row in enumerate(sheet.iter_rows(values_only=True)):
                    yield {"label": f"工作表 {sheet.title} 第 {i + 1} 行", "text": " | ".join(str(v) if v is not None else "" for v in row)}
        finally:
            book.close()
    else:
        from pptx import Presentation
        deck = Presentation(stream)
        for i, slide in enumerate(deck.slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    yield {"label": f"幻灯片 {i + 1}", "text": shape.text}
                if shape.has_table:
                    yield {"label": f"幻灯片 {i + 1} 表格", "text": "\n".join(" | ".join(c.text for c in r.cells) for r in shape.table.rows)}
                if shape.shape_type == 13:
                    yield ({"label": f"幻灯片 {i + 1} 图片", "image": shape.image.blob if raw_images else image_url(shape.image.blob)} if include_images
                           else {"label": f"幻灯片 {i + 1} 图片", "skipped_image": True})
    if suffix != ".pptx":
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            for name in archive.namelist():
                if "/media/" in name and Path(name).suffix.lower() in {".png", ".jpg", ".jpeg", ".gif", ".webp"}:
                    yield ({"label": f"嵌入图片 {Path(name).name}", "image": archive.read(name) if raw_images else image_url(archive.read(name))} if include_images
                           else {"label": f"嵌入图片 {Path(name).name}", "skipped_image": True})
    return


def parse_document(data: bytes, suffix: str):
    return list(iter_document(data, suffix))


@observed('media.resolve_media')
def resolve_media(account, message, max_mb):
    from ..chat_helpers import _resolve_account_dir
    from ..media_helpers import _resolve_media_path_for_kind, _read_and_maybe_decrypt_media, _fallback_search_media_by_file_id, _resolve_account_wxid_dir
    account_dir = _resolve_account_dir(account)
    kind, raw = message["kind"], message["media"]
    md5 = raw.get("imageMd5" if kind == "image" else "fileMd5", "")
    path = _resolve_media_path_for_kind(account_dir, kind=kind, md5=md5, username=message["username"], allow_fallback_scan=False)
    if path is None and kind == "image" and raw.get("imageFileId"):
        path = _fallback_search_media_by_file_id(str(_resolve_account_wxid_dir(account_dir) or ''), raw["imageFileId"], kind="image", username=message["username"], allow_global_scan=False)
    if path is None:
        raise ValueError("本机附件或图片缺失")
    path = Path(path)
    if path.stat().st_size > max_mb * 1024 * 1024:
        raise ValueError(f"附件超过 {max_mb} MB，请调整上限后重试")
    if kind == "image":
        data, _ = _read_and_maybe_decrypt_media(path, account_dir=account_dir)
        return data, ".image"
    return path.read_bytes(), Path(raw.get("title") or path.name).suffix.lower()


class MediaService:
    def __init__(self, store, models):
        self.store, self.models = store, models

    async def enrich(self, account, message, options, vision_profile, checkpoint, unit_callback=None, progress_callback=None):
        if message["kind"] not in {"image", "file"}:
            return message
        return await self._enrich(account, message, options, vision_profile, checkpoint, unit_callback, progress_callback)

    async def _overview(self, account, result, data, suffix, options, profile, checkpoint,
                        key, skip_images, unit_callback, progress_callback):
        cursor = options.get('media_cursor') or {'part_offset': 0, 'text_offset': 0}
        start, text_offset = cursor['part_offset'], cursor['text_offset']
        parts = iter_document(data, suffix, include_images=suffix == '.pdf' and not skip_images, scan_only=True)
        executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix='media-parser')
        fragments, index, consumed, text_bytes, images = [], 0, 0, 0, 0
        completed = 0
        next_cursor, exhausted, skipped = None, False, False

        def progress(phase, label, cached=False):
            if progress_callback:
                progress_callback({'phase': phase, 'label': label, 'completed_units': completed, 'cached': cached})

        try:
            while consumed < 32 and text_bytes < 12288 and images < 3:
                checkpoint()
                progress('extracting', f'附件部件 {index + 1}')
                part = await executor_call(executor, next, parts, None)
                if part is None:
                    exhausted = True
                    break
                if index < start:
                    index += 1
                    continue
                checkpoint()
                label = part['label']
                progress('extracting', label)
                consumed += 1
                if part.get('skipped_image'):
                    skipped = True
                elif 'image' in part:
                    if not profile.get('vision'):
                        raise ValueError('该附件包含扫描图片，请配置视觉模型后重试')
                    partial_key = f'{key}:{index}'
                    partial = self.store.get('media_cache', partial_key)
                    images += 1
                    if unit_callback:
                        unit_callback(label, cached=bool(partial))
                    progress('analyzing', label, cached=bool(partial))
                    if partial:
                        description = partial['text']
                    else:
                        prompt = f'根据这份附件的{label}简要提取标题、用途、主要内容。无法辨认时明确说明，不推断未读页面。'
                        if options.get('question'):
                            prompt += '\n本次问题：' + options['question']
                        description = await self.models.invoke(profile, prompt, images=[part['image']], account=account)
                        if not description.strip():
                            raise ValueError('视觉分析未返回可用内容')
                        self.store.put('media_cache', {'text': description}, id=partial_key, account=account)
                    if not description.strip():
                        raise ValueError('图片分析缓存为空，缓存内容无效')
                    fragments.append(f'[{label}] {description}')
                elif part.get('text', '').strip():
                    full_text = part['text']
                    offset = text_offset if index == start else 0
                    remaining = 12288 - text_bytes
                    fragment = next(pieces(full_text[offset:], remaining), '')
                    # pieces 保留完整字符；剩余空间小于一个字符时留给下一次读取。
                    if size(fragment) > remaining:
                        fragment = ''
                    text_bytes += size(fragment)
                    if fragment:
                        fragments.append(f'[{label}] {fragment}')
                    if offset + len(fragment) < len(full_text):
                        next_cursor = {'part_offset': index, 'text_offset': offset + len(fragment)}
                        completed += 1
                        progress('extracting', label)
                        break
                index += 1
                completed += 1
                progress('extracting', label)
            if not exhausted and next_cursor is None:
                checkpoint()
                # 有界预读只确认是否到末尾；后续请求仍从尚未处理的 index 开始。
                exhausted = await executor_call(executor, next, parts, None) is None
                if not exhausted:
                    next_cursor = {'part_offset': index, 'text_offset': 0}
        finally:
            try:
                await asyncio.shield(executor_call(executor, parts.close))
            finally:
                executor.shutdown(wait=False)
        if not fragments and next_cursor is None:
            raise ValueError('附件未提取到文字，扫描图片未分析，请配置视觉模型后重试' if skipped
                             else '附件未提取到可分析内容')
        detail = {'complete': exhausted and not skipped and start == 0 and text_offset == 0, 'parts_read': consumed, 'text_bytes': text_bytes,
            'images_analyzed': images, 'next_cursor': next_cursor, 'images_not_inspected': skipped,
            'text_only': images == 0, 'start_part': start, 'start_text_offset': text_offset}
        coverage = ('已分析附件概览；已读取标注位置的部分内容，不代表全文分析' if fragments
                    else '附件概览当前片段未提取到可用内容')
        if skipped:
            coverage += '；未检查嵌入图片内容'
        if next_cursor:
            coverage += '；可按返回游标继续读取'
        combined = '\n'.join(fragments)
        if fragments:
            self.store.put('media_cache', {'text': combined, 'units': images, 'coverage': coverage,
                'analysis_mode': 'overview', 'coverage_detail': detail}, id=key, account=account)
        progress('completed', '附件概览完成')
        return result | {'text': result['text'] + '\n' + combined, 'coverage': coverage,
            'analysis_mode': 'overview', 'coverage_detail': detail}

    @observed('media.enrich')
    async def _enrich(self, account, message, options, vision_profile, checkpoint, unit_callback=None, progress_callback=None):
        result = dict(message)
        mode = options.get('analysis_mode', 'full') if message['kind'] == 'file' else 'full'
        if mode not in ('overview', 'full'):
            raise ValueError('附件分析模式必须是 overview 或 full')
        if not options.get("media", True):
            return result | {"coverage": "未启用媒体分析"}
        skip_images = bool(options.get('skip_unsupported_images') and not vision_profile.get('vision'))
        skipped_notice = '当前模型不支持图片，已跳过图片内容'
        if skip_images and message['kind'] == 'image':
            return result | {'coverage': skipped_notice}
        parts = None
        executor = None
        try:
            if progress_callback:
                progress_callback({'phase': 'extracting', 'label': '定位附件', 'completed_units': 0})
            data, suffix = await asyncio.to_thread(resolve_media, account, message, options.get("max_attachment_mb", 20))
            diagnostic_event('media.located', suffix=suffix, bytes=len(data))
            signature = {"version": 2, "profile": vision_profile.get("id"), "model": vision_profile.get('model'),
                         "revision": vision_profile.get("revision"), "suffix": suffix, 'skip_images': skip_images}
            if mode == 'overview':
                signature.update(analysis_mode=mode, media_cursor=options.get('media_cursor'), overview_version=1)
            question = str(options.get('question') or '').strip()
            if question:
                signature['question'] = question
            key = hashlib.sha256(account.encode() + data + json.dumps(signature, sort_keys=True).encode()).hexdigest()
            cached = self.store.get("media_cache", key)
            if cached:
                if not cached['text'].strip():
                    raise ValueError('附件分析缓存为空，缓存内容无效')
                diagnostic_event('media.cache', cached=True, count=cached.get('units', 1))
                if unit_callback:
                    for n in range(cached.get('units', 1)):
                        unit_callback(f'缓存图片 {n + 1}', cached=True)
                if progress_callback:
                    progress_callback({'phase': 'completed', 'label': '附件分析缓存',
                        'completed_units': cached.get('coverage_detail', {}).get('parts_read', cached.get('parts_read', cached.get('units', 1))), 'cached': True})
                return result | {"text": result["text"] + "\n" + cached["text"], "coverage": cached.get('coverage', '已分析') + '（缓存）',
                    **{k: cached[k] for k in ('analysis_mode', 'coverage_detail') if k in cached}}
            if mode == 'overview':
                return await self._overview(account, result, data, suffix, options, vision_profile, checkpoint,
                    key, skip_images, unit_callback, progress_callback)
            parts = iter([{ "label": "图片", "image": image_url(data)}]) if suffix == ".image" else iter_document(data, suffix, include_images=not skip_images)
            executor = ThreadPoolExecutor(max_workers=1, thread_name_prefix='media-parser')
            text, index, units = [], -1, 0
            local_text = []
            skipped, has_text = False, False
            while True:
                checkpoint()
                unit_started = time.monotonic()
                diagnostic_event('media.page.started', index=index+1)
                if progress_callback:
                    progress_callback({'phase': 'extracting', 'label': f'附件部件 {index + 2}', 'completed_units': index + 1})
                part = await executor_call(executor, next, parts, None)
                if part is None:
                    diagnostic_event('media.page.finished', index=index+1, complete=True)
                    break
                index += 1
                checkpoint()
                label = part["label"]
                if part.get('skipped_image'):
                    skipped = True
                    continue
                if "image" in part:
                    if not vision_profile.get("vision"):
                        raise ValueError("该附件包含图片，请配置视觉模型后重试")
                    partial_key = f"{key}:{index}"
                    partial = self.store.get("media_cache", partial_key)
                    units += 1
                    if unit_callback:
                        unit_callback(label, cached=bool(partial))
                    if partial:
                        diagnostic_event('media.page.cache', index=index, cached=True)
                        description = partial["text"]
                    else:
                        if progress_callback:
                            progress_callback({'phase': 'analyzing', 'label': label, 'completed_units': index})
                        prompt = f"描述这份聊天资料中的{label}，完整提取可辨认文字、表格和关键信息。不能辨认的内容请说明。"
                        if question:
                            prompt += '\n重点核查本次问题：' + question
                        description = await self.models.invoke(vision_profile, prompt, images=[part["image"]], account=account)
                        if not description.strip():
                            raise ValueError('视觉分析未返回可用内容')
                        self.store.put("media_cache", {"text": description}, id=partial_key, account=account)
                    if not description.strip():
                        raise ValueError('图片分析缓存为空，缓存内容无效')
                    text.append(f"[{label}] {description}")
                else:
                    has_text = has_text or bool(part.get('text', '').strip())
                    text.append(f"[{label}] {part.get('text', '')}")
                    local_text.append(text[-1])
                    # 仅记录本地提取文字，独立于视觉模型输出，供离线检索复用。
                    local_id = hashlib.sha256(f"{account}:{message['username']}:{message['anchor']}".encode()).hexdigest()
                    file_hash = hashlib.sha256(data).hexdigest()
                    with self.store.lock:
                        previous = self.store.get('local_media_text', local_id)
                        # Explicit extraction has already read this exact file.
                        # Interrupted Agent progress must not replace its complete text.
                        if not (previous is not None and previous.get('complete') is True and previous.get('file_hash') == file_hash):
                            self.store.put('local_media_text', {'text': '\n'.join(local_text), 'username': message['username'],
                                'anchor': message['anchor'], 'file_hash': file_hash, 'version': 1}, id=local_id, account=account)
                diagnostic_event('media.page.finished', index=index, duration_ms=(time.monotonic()-unit_started)*1000)
                if progress_callback:
                    progress_callback({'phase': 'extracting', 'label': label, 'completed_units': index + 1})
            if not has_text and not units and not skipped:
                raise ValueError('附件未提取到可分析内容')
            combined = "\n".join(text)
            coverage = ('已分析附件文字；' + skipped_notice) if skipped and has_text else (skipped_notice if skipped else '已分析')
            self.store.put("media_cache", {"text": combined, "units": units, 'parts_read': index + 1, 'coverage': coverage}, id=key, account=account)
            if progress_callback:
                progress_callback({'phase': 'completed', 'label': '附件分析完成', 'completed_units': index + 1})
            return result | {"text": result["text"] + "\n" + combined, "coverage": coverage}
        except (ValueError, OSError, zipfile.BadZipFile, ElementTree.ParseError, KeyError) as exc:
            diagnostic_event('media.failed', level=logging.WARNING, error=exc, reason_code=media_failure_reason(exc))
            return result | {"coverage": str(exc), "text": result["text"] + "\n[附件未分析]"}
        finally:
            if executor is not None:
                try:
                    if hasattr(parts, 'close'):
                        await asyncio.shield(executor_call(executor, parts.close))
                finally:
                    executor.shutdown(wait=False)


def media_failure_reason(error):
    # 仅将本地错误映射到固定分类，不记录异常正文。
    text = str(error)
    for needle, code in [('大小', 'size_limit'), ('超过', 'size_limit'), ('加密', 'encrypted'), ('损坏', 'corrupt'), ('不存在', 'missing'), ('找到', 'missing'), ('视觉', 'vision_required'), ('支持', 'unsupported')]:
        if needle in text:
            return code
    return 'parse_failed'
